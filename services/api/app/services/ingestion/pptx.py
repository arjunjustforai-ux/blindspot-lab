"""PPTX extraction with python-pptx (+ LibreOffice previews).

Slide-anchored units (slide numbers are 1-based and preserved exactly):
  - slide_text   title + body text from every text-bearing shape (groups included)
  - slide_table  native tables serialised as Markdown
  - slide_chart  native chart data (type, categories, series values) — exact numbers
                 from the chart XML, not OCR
  - slide_figure embedded pictures: OCR text (+ optional vision description)
  - slide_notes  speaker notes
Previews: each slide rendered to PNG through LibreOffice headless -> PDF -> PyMuPDF.
"""

from __future__ import annotations

import shutil
import subprocess
import tempfile
import uuid
from pathlib import Path

import pymupdf
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.exc import PackageNotFoundError

from app.config import get_settings
from app.services.ai import vision
from app.services.ai.text import looks_like_injection, normalize_ws
from app.services.ingestion.types import ExtractedUnit, ExtractionError, ExtractionResult, ProgressFn


def _iter_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(sh.shapes)
        else:
            yield sh


def _fmt_num(v: object) -> str:
    if isinstance(v, float):
        return f"{v:g}"
    return str(v)


def render_previews(path: Path, expected: int) -> tuple[dict[int, bytes], str | None]:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return {}, "LibreOffice not installed; slide previews unavailable"
    with tempfile.TemporaryDirectory() as tmp:
        try:
            subprocess.run(
                [soffice, "--headless", "--norestore", f"-env:UserInstallation=file://{tmp}/lo",
                 "--convert-to", "pdf", "--outdir", tmp, str(path)],
                check=True, capture_output=True, timeout=240,
            )
        except (subprocess.SubprocessError, OSError) as e:
            return {}, f"Slide preview rendering failed: {e}"
        pdfs = list(Path(tmp).glob("*.pdf"))
        if not pdfs:
            return {}, "Slide preview rendering produced no output"
        doc = pymupdf.open(pdfs[0])
        if doc.page_count != expected:
            return {}, f"Preview page count {doc.page_count} != slide count {expected}; previews disabled"
        out = {}
        for page in doc:
            zoom = 1280 / page.rect.width
            out[page.number + 1] = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom)).tobytes("png")
        return out, None


def extract_pptx(path: Path, *, course_id: uuid.UUID | None, progress: ProgressFn) -> ExtractionResult:
    s = get_settings()
    try:
        prs = Presentation(str(path))
    except (PackageNotFoundError, KeyError, ValueError) as e:
        raise ExtractionError(f"Could not open PPTX (is it a valid PowerPoint file?): {e}") from e

    n = len(prs.slides)
    if n == 0:
        raise ExtractionError("Presentation has no slides")
    deck_title = normalize_ws(prs.core_properties.title or "") or path.stem
    units: list[ExtractedUnit] = []
    warnings: list[str] = []

    progress("Rendering slide previews", 0.05)
    previews, warn = render_previews(path, n)
    if warn:
        warnings.append(warn)

    deck_text = " ".join(
        sh.text_frame.text for sl in prs.slides for sh in _iter_shapes(sl.shapes) if getattr(sh, "has_text_frame", False)
    )
    currency = next((c for c in ("₹", "€", "£") if c in deck_text), None)
    vision_calls = 0
    for idx, slide in enumerate(prs.slides, start=1):
        progress(f"Reading slide {idx}/{n}", 0.15 + 0.7 * (idx - 1) / n)
        title_shape = slide.shapes.title
        title = normalize_ws(title_shape.text_frame.text) if title_shape is not None and title_shape.has_text_frame else ""
        heading = title or f"Slide {idx}"
        sp = [deck_title, heading]
        body_lines: list[str] = []
        picture_count = 0
        shape_count = 0
        for sh in _iter_shapes(slide.shapes):
            shape_count += 1
            if title_shape is not None and sh.shape_id == title_shape.shape_id:
                continue
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    t = normalize_ws("".join(r.text for r in p.runs))
                    if t:
                        body_lines.append(("  " * p.level) + "- " + t if p.level else t)
            if getattr(sh, "has_table", False) and sh.has_table:
                rows = [[normalize_ws(c.text) for c in r.cells] for r in sh.table.rows]
                if rows:
                    md = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])]
                    md += ["| " + " | ".join(r) + " |" for r in rows[1:]]
                    units.append(ExtractedUnit(
                        modality="slide_table", content=f"Table on slide {idx} ({heading}):\n" + "\n".join(md),
                        heading=heading, section_path=sp, slide_number=idx,
                        extraction_meta={"rows": len(rows), "cols": len(rows[0])},
                    ))
            if getattr(sh, "has_chart", False) and sh.has_chart:
                ch = sh.chart
                try:
                    cats = [str(c) for c in ch.plots[0].categories]
                except (IndexError, AttributeError):
                    cats = []
                series_txt = []
                for plot in ch.plots:
                    for ser in plot.series:
                        vals = list(ser.values)
                        pairs = "; ".join(f"{c}: {_fmt_num(v)}" for c, v in zip(cats, vals)) if cats else \
                            ", ".join(_fmt_num(v) for v in vals)
                        series_txt.append(f"Series '{ser.name}': {pairs}")
                ctitle = ""
                try:
                    if ch.has_title:
                        ctitle = normalize_ws(ch.chart_title.text_frame.text)
                except AttributeError:
                    pass
                ctype = str(ch.chart_type).split(".")[-1].split(" ")[0].lower()
                units.append(ExtractedUnit(
                    modality="slide_chart",
                    content=f"Chart on slide {idx} ({heading}){': ' + ctitle if ctitle else ''} [{ctype}]\n"
                    + "\n".join(series_txt),
                    heading=heading, section_path=sp, slide_number=idx,
                    extraction_meta={"chart_type": ctype, "categories": cats, "method": "native-chart-xml"},
                ))
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                picture_count += 1
                try:
                    blob = sh.image.blob
                    img = vision.image_from_bytes(blob)
                except Exception:
                    continue
                if not vision.is_meaningful_image(img):
                    continue
                png = vision.png_bytes(img)
                ocr = vision.ocr_image(img, currency_symbol=currency)
                content = f"[Picture on slide {idx}: {heading}]"
                if ocr.word_count:
                    content += "\nText and labels read from the picture (OCR): " + "; ".join(ocr.lines)
                units.append(ExtractedUnit(
                    modality="slide_figure", content=content, heading=heading, section_path=sp, slide_number=idx,
                    image_png=png, extraction_meta={"ocr_words": ocr.word_count,
                                                    "ocr_mean_confidence": ocr.mean_confidence,
                                                    "method": "tesseract-ocr"},
                ))
        text = "\n".join([title] + body_lines) if title else "\n".join(body_lines)
        if text.strip():
            meta = {"shape_count": shape_count, "pictures": picture_count}
            if looks_like_injection(text):
                meta["injection_suspect"] = True
            units.append(ExtractedUnit(
                modality="slide_text", content=text, heading=heading, section_path=sp, slide_number=idx,
                extraction_meta=meta,
            ))
        if slide.has_notes_slide:
            notes = normalize_ws(slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame else "")
            if notes:
                units.append(ExtractedUnit(
                    modality="slide_notes", content=f"Speaker notes for slide {idx} ({heading}): {notes}",
                    heading=heading, section_path=sp, slide_number=idx,
                    extraction_meta={"injection_suspect": True} if looks_like_injection(notes) else {},
                ))
        # Vision model on the rendered slide when it carries visual content.
        if (picture_count or shape_count >= 4) and idx in previews and vision_calls < s.max_vision_calls_per_source:
            desc, vinfo = vision.describe_figure(previews[idx], context=heading, course_id=course_id)
            if desc is not None:
                vision_calls += 1
                if not desc.get("is_decorative"):
                    units.append(ExtractedUnit(
                        modality="slide_figure",
                        content=f"Model description of slide {idx} ({heading}): {desc.get('description', '')}"
                        + (f"\nTakeaway: {desc.get('key_takeaway')}" if desc.get("key_takeaway") else ""),
                        heading=heading, section_path=sp, slide_number=idx, image_png=previews[idx],
                        is_generated=True, generation_meta={"kind": "vision_description", **vinfo, "raw": desc},
                    ))

    if not units:
        raise ExtractionError("No text, notes, tables, charts or pictures could be extracted from this deck")
    return ExtractionResult(
        units=units, slide_count=n, previews=previews, warnings=warnings,
        metadata={"title": deck_title, "author": prs.core_properties.author or None, "vision_calls": vision_calls,
                  "previews": bool(previews)},
    )
