"""Build the original, self-contained Blindspot Lab demonstration course pack.

Outputs (in demo/course-materials/):
  - Marketing_Analytics_Primer.pdf   textbook-style chapter with figures and tables
  - Unit_Economics_for_Marketers.pptx slide deck with a native chart, a native table,
                                       a shape-based funnel diagram, a picture diagram
                                       and speaker notes
  - Lecture_03_Price_Cuts.mp4         narrated lecture video (synthesized speech via
                                       Piper TTS) with on-screen frames
  - lecture_03_narration.json          the narration script and the planned segment
                                       timings (ground truth for evaluation only; the
                                       app never reads this file)

All companies, prices and results are fictional and explicitly synthetic.

Requirements (dev-only, not needed by the app):
  pip install reportlab matplotlib python-pptx pillow piper-tts
  ffmpeg on PATH, and a Piper voice model (see --voice).

Usage:
  python demo/scripts/build_course_pack.py --voice ~/.cache/piper/en_US-lessac-medium.onnx
"""

from __future__ import annotations

import argparse
import json
import subprocess
import tempfile
import wave
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from PIL import Image, ImageDraw, ImageFont  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "demo" / "course-materials"
# DejaVu fonts bundled with matplotlib (portable; include the rupee sign).
FONT_DIR = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
SYNTHETIC_NOTE = (
    "Synthetic teaching material created for the Blindspot Lab demonstration. "
    "All companies, prices and results are fictional."
)

INDIGO = "#4f46e5"
CYAN = "#0891b2"
AMBER = "#d97706"
SLATE = "#334155"
GREEN = "#059669"
RED = "#dc2626"


# ---------------------------------------------------------------------------
# Figures (rendered with matplotlib; DejaVu Sans supports the rupee sign)
# ---------------------------------------------------------------------------


def fig_price_cut(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(6.4, 3.4), dpi=200)
    labels = ["Preserve revenue", "Preserve contribution"]
    values = [11.1, 33.3]
    bars = ax.bar(labels, values, color=[CYAN, INDIGO], width=0.55)
    for b, v in zip(bars, values):
        ax.text(b.get_x() + b.get_width() / 2, v + 1, f"+{v:.1f}% volume", ha="center", fontsize=12, weight="bold")
    ax.set_ylim(0, 42)
    ax.set_ylabel("Volume increase needed (%)", fontsize=10)
    ax.set_title("Price cut from ₹500 to ₹450 (variable cost ₹300)", fontsize=12)
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def fig_roas_channels(path: Path) -> None:
    channels = ["Search", "Social", "Display", "Email"]
    roas = [4.0, 2.0, 1.5, 9.0]
    contrib = [60, -16, -20, 26]  # thousand rupees, CM ratio 40%
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(8.0, 3.6), dpi=200)
    b1 = a1.bar(channels, roas, color=INDIGO)
    a1.axhline(2.5, color=AMBER, linestyle="--", linewidth=1.5)
    a1.text(-0.45, 9.9, "Break-even ROAS 2.5x (dashed line)", color=AMBER, fontsize=9, ha="left")
    for b, v in zip(b1, roas):
        a1.text(b.get_x() + b.get_width() / 2, v + 0.35, f"{v:.1f}x", ha="center", fontsize=11, weight="bold",
                bbox=dict(facecolor="white", edgecolor="none", pad=1.5), zorder=5)
    a1.set_title("ROAS by channel", fontsize=12)
    a1.set_ylim(0, 10.5)
    a1.spines[["top", "right"]].set_visible(False)
    colors = [GREEN if c > 0 else RED for c in contrib]
    b2 = a2.bar(channels, contrib, color=colors)
    a2.axhline(0, color=SLATE, linewidth=1)
    for b, v in zip(b2, contrib):
        label = f"+₹{v}k" if v > 0 else f"-₹{abs(v)}k"
        y = v + (3 if v > 0 else -7)
        a2.text(b.get_x() + b.get_width() / 2, y, label, ha="center", fontsize=11, weight="bold")
    a2.set_title("Contribution after ad spend", fontsize=12)
    a2.set_ylim(-30, 75)
    a2.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def fig_order_values(path: Path) -> None:
    values = [800, 900, 1000, 1100, 9200]
    fig, ax = plt.subplots(figsize=(6.4, 3.0), dpi=200)
    ax.bar([f"Order {i + 1}" for i in range(5)], values, color=CYAN)
    ax.axhline(2600, color=AMBER, linestyle="--")
    ax.axhline(1000, color=INDIGO, linestyle=":")
    ax.text(4.45, 2800, "mean ₹2,600", color=AMBER, ha="right", fontsize=10, weight="bold")
    ax.text(4.45, 1250, "median ₹1,000", color=INDIGO, ha="right", fontsize=10, weight="bold")
    ax.set_title("Five synthetic order values", fontsize=12)
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def fig_waterfall(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(7.5, 3.8), dpi=200)
    steps = ["Discounted price", "Variable cost", "Contribution"]
    ax.bar(steps[0], 1080, color=INDIGO)
    ax.bar(steps[1], -800, bottom=1080, color=RED)
    ax.bar(steps[2], 280, color=GREEN)
    ax.text(0, 1110, "₹1,080", ha="center", fontsize=12, weight="bold")
    ax.text(1, 680, "-₹800", ha="center", fontsize=12, weight="bold", color="white")
    ax.text(2, 310, "₹280", ha="center", fontsize=12, weight="bold")
    ax.set_title("Where the discounted jacket price goes (per unit)", fontsize=12)
    ax.set_ylim(0, 1250)
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


# ---------------------------------------------------------------------------
# PDF chapter (reportlab)
# ---------------------------------------------------------------------------


def build_pdf(path: Path, figdir: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        Image as RLImage,
    )
    from reportlab.platypus import (
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    pdfmetrics.registerFont(TTFont("DejaVu", str(FONT_DIR / "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(FONT_DIR / "DejaVuSans-Bold.ttf")))
    pdfmetrics.registerFont(TTFont("DejaVu-Oblique", str(FONT_DIR / "DejaVuSans-Oblique.ttf")))

    body = ParagraphStyle("body", fontName="DejaVu", fontSize=10.5, leading=15, spaceAfter=7)
    h1 = ParagraphStyle("h1", fontName="DejaVu-Bold", fontSize=18, leading=23, spaceAfter=10, spaceBefore=4)
    h2 = ParagraphStyle("h2", fontName="DejaVu-Bold", fontSize=13, leading=17, spaceAfter=6, spaceBefore=8)
    caption = ParagraphStyle("caption", fontName="DejaVu-Oblique", fontSize=9, leading=12, spaceAfter=10)
    title = ParagraphStyle("title", fontName="DejaVu-Bold", fontSize=24, leading=30, alignment=TA_CENTER)
    subtitle = ParagraphStyle("subtitle", fontName="DejaVu", fontSize=13, leading=18, alignment=TA_CENTER)
    note = ParagraphStyle("note", fontName="DejaVu-Oblique", fontSize=9.5, leading=13, textColor=colors.HexColor("#475569"))
    formula = ParagraphStyle(
        "formula", fontName="DejaVu-Bold", fontSize=10.5, leading=15, leftIndent=18, spaceAfter=8,
        textColor=colors.HexColor("#312e81"),
    )

    def table(data: list[list[str]], widths: list[float]) -> Table:
        t = Table(data, colWidths=widths)
        t.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, -1), "DejaVu"),
                    ("FONTNAME", (0, 0), (-1, 0), "DejaVu-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#94a3b8")),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e0e7ff")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        return t

    P = Paragraph
    s: list = []
    # Page 1 — cover
    s += [
        Spacer(1, 3 * cm),
        P("Making Better Decisions with Data", title),
        Spacer(1, 0.4 * cm),
        P("A Primer on Marketing and Business Analytics", subtitle),
        Spacer(1, 0.3 * cm),
        P("Blindspot Lab Demonstration Unit · Chapter Pack", subtitle),
        Spacer(1, 2 * cm),
        P("About this unit", h2),
        P(
            "This short primer introduces six connected ideas that marketing and business analysts use to make "
            "decisions: revenue and conversion metrics, contribution margin, return on advertising spend, customer "
            "acquisition cost and lifetime value, correlation versus causation, and decision-making under changing "
            "assumptions. Each chapter builds on the previous ones.",
            body,
        ),
        P(SYNTHETIC_NOTE, note),
        PageBreak(),
    ]

    # Page 2 — Chapter 1
    s += [
        P("Chapter 1. Revenue and Conversion Metrics", h1),
        P("1.1 Revenue", h2),
        P(
            "Revenue is the money a business takes in from sales before any costs are deducted. For a single product, "
            "revenue equals the selling price multiplied by the number of units sold. Revenue is a measure of scale, "
            "not of profit: a business can grow revenue while losing money on every sale.",
            body,
        ),
        P("Revenue = selling price × units sold", formula),
        P("1.2 Conversion rate and absolute conversions", h2),
        P(
            "The conversion rate is the share of visitors (or leads) who complete a desired action such as a "
            "purchase. A high conversion rate does not guarantee a large number of conversions. A landing page that "
            "converts 8% of 500 visitors produces 40 customers, while a page that converts 3% of 4,000 visitors "
            "produces 120 customers. Decisions about where to invest should consider absolute conversions and the "
            "value of each conversion, not only the rate.",
            body,
        ),
        P("Conversion rate = conversions ÷ visitors", formula),
        P("Table 1.1 Synthetic funnel data for Northwind Outdoor Co. (fictional), one month", caption),
        table(
            [
                ["Channel", "Visitors", "Conversion rate", "Conversions", "Avg order value", "Revenue"],
                ["Email", "5,000", "4.0%", "200", "₹2,500", "₹500,000"],
                ["Paid search", "20,000", "2.5%", "500", "₹2,200", "₹1,100,000"],
                ["Social", "40,000", "0.8%", "320", "₹1,800", "₹576,000"],
                ["Affiliates", "8,000", "3.0%", "240", "₹2,000", "₹480,000"],
            ],
            [2.4 * cm, 2.0 * cm, 3.2 * cm, 2.5 * cm, 3.1 * cm, 3.0 * cm],
        ),
        Spacer(1, 0.3 * cm),
        P(
            "In Table 1.1, email has the highest conversion rate, but paid search produces the most conversions and "
            "the most revenue because it reaches many more visitors.",
            body,
        ),
        P("1.3 Percentage points versus percent change", h2),
        P(
            "When a rate changes, there are two different ways to describe the change. A percentage-point change is "
            "the simple difference between two percentages: a conversion rate that rises from 2% to 3% has increased "
            "by 1 percentage point. A percent change is the difference relative to the starting value: the same move "
            "from 2% to 3% is a 50% increase. Reports should state which measure they use, because confusing the two "
            "can make a small change look large or a large change look small.",
            body,
        ),
        PageBreak(),
    ]

    # Pages 3-4 — Chapter 2
    s += [
        P("Chapter 2. Contribution Margin", h1),
        P("2.1 Variable costs and fixed costs", h2),
        P(
            "Variable costs change in proportion to the number of units sold, for example materials, packaging, "
            "payment fees and delivery. Fixed costs, such as rent and salaried staff, do not change with short-run "
            "volume. Separating the two is the first step in understanding how a decision affects profit.",
            body,
        ),
        P("2.2 Unit contribution and the contribution margin ratio", h2),
        P(
            "Contribution margin per unit is the selling price minus the variable cost per unit. It is the amount each "
            "additional unit contributes towards covering fixed costs and then generating profit. Total contribution "
            "equals the contribution margin per unit multiplied by the number of units sold. The contribution margin "
            "ratio expresses contribution as a share of the selling price.",
            body,
        ),
        P("Contribution margin per unit = price − variable cost per unit", formula),
        P("Contribution margin ratio = (price − variable cost) ÷ price", formula),
        P(
            "Example: a product sells for ₹500 and has a variable cost of ₹300 per unit. Its contribution margin per "
            "unit is ₹200 and its contribution margin ratio is 40%.",
            body,
        ),
        P("2.3 Price changes and the volume needed to compensate", h2),
        P(
            "A price cut reduces the contribution of every unit sold, not just the revenue. Because variable costs "
            "stay the same, the contribution per unit falls by the full amount of the price cut. To keep total "
            "contribution unchanged, volume must rise by the ratio of the old to the new contribution per unit.",
            body,
        ),
        P("Required volume increase = (old contribution per unit ÷ new contribution per unit) − 1", formula),
        P(
            "Continuing the example, cutting the price from ₹500 to ₹450 lowers contribution per unit from ₹200 to "
            "₹150. Volume must rise by 200 ÷ 150 − 1 = 33.3% just to earn the same total contribution. Keeping "
            "revenue constant would need only an 11.1% increase (500 ÷ 450 − 1), which is why plans that aim to "
            "'make up the revenue' can still leave the business with less contribution to cover its fixed costs.",
            body,
        ),
        P(
            "Preserving revenue and preserving contribution are different goals. Revenue preservation ignores "
            "variable costs; contribution preservation accounts for them.",
            body,
        ),
        RLImage(str(figdir / "fig2_1.png"), width=13.5 * cm, height=13.5 * cm * 3.4 / 6.4),
        P(
            "Figure 2.1 Volume increase required after a 10% price cut (synthetic example). Preserving contribution "
            "requires three times the volume increase needed to preserve revenue.",
            caption,
        ),
        P("2.4 Total cost versus marginal cost", h2),
        P(
            "For decisions about one additional order, what matters is the marginal (incremental) cost of serving "
            "it, usually the variable cost, not the average total cost per unit, which includes fixed costs that will "
            "be paid anyway. Example: a workshop with a variable cost of ₹300 and an average total cost of ₹380 per "
            "unit (including allocated rent) receives a one-off order for 100 units at ₹340 each. If spare capacity "
            "exists and fixed costs will not change, the order adds (340 − 300) × 100 = ₹4,000 of contribution, even "
            "though ₹340 is below the average total cost.",
            body,
        ),
        P("2.5 Break-even volume", h2),
        P(
            "Break-even volume is the number of units at which total contribution exactly covers fixed costs. With "
            "fixed costs of ₹60,000 and contribution per unit of ₹200, the business must sell 300 units to break even.",
            body,
        ),
        P("Break-even volume = fixed costs ÷ contribution margin per unit", formula),
        PageBreak(),
    ]

    # Chapter 3
    s += [
        P("Chapter 3. Return on Advertising Spend", h1),
        P("3.1 Definition", h2),
        P(
            "Return on advertising spend (ROAS) is the revenue attributed to advertising divided by the advertising "
            "cost. A campaign that spends ₹50,000 and is credited with ₹200,000 of revenue has a ROAS of 4, often "
            "written 4x or 400%.",
            body,
        ),
        P("ROAS = attributed revenue ÷ advertising spend", formula),
        P("3.2 ROAS is not profitability", h2),
        P(
            "ROAS compares advertising cost with revenue, not with profit. Because revenue still has to pay for the "
            "variable costs of the products sold, a ROAS above 1 does not mean the campaign made money. The "
            "break-even ROAS is the ROAS at which the contribution from attributed sales exactly covers the "
            "advertising spend. A business with a 40% contribution margin ratio needs a ROAS of at least 2.5 before "
            "advertising contributes anything towards fixed costs.",
            body,
        ),
        P("Break-even ROAS = 1 ÷ contribution margin ratio", formula),
        P("Contribution after advertising = attributed revenue × contribution margin ratio − advertising spend", formula),
        RLImage(str(figdir / "fig3_1.png"), width=15.5 * cm, height=15.5 * cm * 3.6 / 8.0),
        P(
            "Figure 3.1 ROAS and contribution after advertising by channel for Northwind Outdoor Co. (synthetic "
            "data, contribution margin ratio 40%, break-even ROAS 2.5).",
            caption,
        ),
        P("3.3 Attribution and incrementality", h2),
        P(
            "ROAS depends on how revenue is attributed to campaigns. Some attributed sales would have happened "
            "without the advertising, so ROAS can overstate the incremental effect of a campaign. Incrementality "
            "tests, which compare exposed and unexposed groups, give a better estimate of the causal effect.",
            body,
        ),
        PageBreak(),
    ]

    # Chapter 4
    s += [
        P("Chapter 4. Customer Acquisition Cost and Lifetime Value", h1),
        P("4.1 Customer acquisition cost", h2),
        P(
            "Customer acquisition cost (CAC) is the total sales and marketing spend in a period divided by the number "
            "of new customers acquired in that period. Spending ₹300,000 to acquire 1,200 new customers gives a CAC "
            "of ₹250.",
            body,
        ),
        P("CAC = sales and marketing spend ÷ new customers acquired", formula),
        P("4.2 Customer lifetime value", h2),
        P(
            "Customer lifetime value (CLV) estimates the total contribution a customer generates over their "
            "relationship with the business. A simple version multiplies the average monthly contribution per "
            "customer by the expected customer lifetime in months. Using revenue instead of contribution overstates "
            "CLV, because it ignores the cost of serving the customer. For example, a customer who pays ₹500 per "
            "month at a 40% contribution margin ratio generates ₹200 of contribution per month; with an expected "
            "lifetime of 18 months, CLV is ₹3,600.",
            body,
        ),
        P("CLV = average monthly contribution per customer × expected lifetime in months", formula),
        P("4.3 Comparing CAC with CLV", h2),
        P(
            "A CLV-to-CAC ratio compares the value a customer generates with the cost of acquiring them. Many "
            "analysts look for a ratio comfortably above 1, often around 3, but the right threshold depends on the "
            "business. The payback period is the number of months of contribution needed to recover CAC: with a CAC "
            "of ₹900 and monthly contribution of ₹200, payback takes 4.5 months.",
            body,
        ),
        P("Payback period (months) = CAC ÷ monthly contribution per customer", formula),
        P(
            "A low CAC is not automatically good. Cheap customers who churn quickly can be worth less than more "
            "expensive customers who stay for years. Always compare CAC with contribution-based lifetime value.",
            body,
        ),
        PageBreak(),
    ]

    # Chapter 5
    s += [
        P("Chapter 5. Correlation, Causation and Experiments", h1),
        P("5.1 Correlation is not causation", h2),
        P(
            "Two variables are correlated when they tend to move together. Correlation alone does not show that one "
            "causes the other. Customers who use a loyalty app may spend more than other customers, but loyal, "
            "high-spending customers may simply be more likely to install the app in the first place. A third "
            "variable that influences both, called a confounder, can create a correlation without any causal link.",
            body,
        ),
        P("5.2 Randomised experiments", h2),
        P(
            "A randomised controlled experiment, often called an A/B test, assigns customers to groups at random so "
            "that confounders are balanced on average. If the treated group then behaves differently from the "
            "control group, the difference can be attributed to the treatment, within the limits of sampling error.",
            body,
        ),
        P("5.3 Statistical significance versus practical significance", h2),
        P(
            "A statistically significant result is one that would be unlikely if there were truly no effect; it says "
            "little about whether the effect is large enough to matter. With very large samples, a 0.1 "
            "percentage-point lift in conversion can be statistically significant yet too small to justify the cost "
            "of a change. Practical significance asks whether the size of the effect is worth acting on, considering "
            "costs, risks and the confidence interval.",
            body,
        ),
        PageBreak(),
    ]

    # Chapter 6
    s += [
        P("Chapter 6. Decisions under Changing Assumptions", h1),
        P("6.1 Sensitivity and scenario analysis", h2),
        P(
            "Business cases depend on assumptions such as price, variable cost, conversion rate and churn. "
            "Sensitivity analysis changes one assumption at a time to see how much the result moves. Scenario "
            "analysis changes several assumptions together to describe a pessimistic, expected and optimistic case.",
            body,
        ),
        P("Table 6.1 Campaign scenarios for 50,000 visitors and ₹100,000 of advertising spend (synthetic)", caption),
        table(
            [
                ["Assumption", "Pessimistic", "Expected", "Optimistic"],
                ["Conversion rate", "1.5%", "2.0%", "2.5%"],
                ["Average order value", "₹1,800", "₹2,000", "₹2,200"],
                ["Contribution margin ratio", "35%", "40%", "45%"],
                ["Contribution after ad spend", "₹372,500", "₹700,000", "₹1,137,500"],
            ],
            [5.0 * cm, 3.2 * cm, 3.2 * cm, 3.2 * cm],
        ),
        Spacer(1, 0.3 * cm),
        P("6.2 Mean versus median", h2),
        P(
            "Averages can mislead when data are skewed. Order values often contain a few very large orders. For the "
            "order values ₹800, ₹900, ₹1,000, ₹1,100 and ₹9,200, the mean is ₹2,600 but the median is ₹1,000. The "
            "median better describes a typical order, while the mean is pulled up by the single large order. Use the "
            "median to describe a typical customer and the mean when you need totals, because the mean multiplied by "
            "the count gives the total.",
            body,
        ),
        RLImage(str(figdir / "fig6_1.png"), width=12.5 * cm, height=12.5 * cm * 3.0 / 6.4),
        P("Figure 6.1 One very large order pulls the mean far above the median (synthetic data).", caption),
        P("6.3 A decision checklist", h2),
        P(
            "First, state the decision and the metric that matters, which is usually contribution rather than "
            "revenue. Second, separate correlation from causal evidence. Third, test how sensitive the conclusion is "
            "to the key assumptions. Fourth, prefer incremental measures over attributed or average measures.",
            body,
        ),
        PageBreak(),
    ]

    # Glossary
    s += [
        P("Glossary", h1),
        P("<b>Break-even ROAS</b>: the ROAS at which contribution from attributed sales equals advertising spend; "
          "1 ÷ contribution margin ratio.", body),
        P("<b>Contribution margin</b>: selling price minus variable cost, per unit or in total.", body),
        P("<b>Confounder</b>: a variable that influences both a supposed cause and its effect.", body),
        P("<b>Conversion rate</b>: conversions divided by visitors or leads.", body),
        P("<b>Customer acquisition cost (CAC)</b>: sales and marketing spend divided by new customers acquired.", body),
        P("<b>Customer lifetime value (CLV)</b>: expected total contribution from a customer relationship.", body),
        P("<b>Percentage point</b>: the arithmetic difference between two percentages.", body),
        P("<b>Variable cost</b>: a cost that rises in proportion to units sold.", body),
        Spacer(1, 0.6 * cm),
        P(SYNTHETIC_NOTE, note),
    ]

    doc = SimpleDocTemplate(
        str(path), pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
        title="Making Better Decisions with Data", author="Blindspot Lab team (synthetic demo material)",
        subject="Marketing and Business Analytics primer",
    )

    def footer(canvas, doc_):  # noqa: ANN001
        canvas.saveState()
        canvas.setFont("DejaVu", 8)
        canvas.setFillColor(colors.HexColor("#64748b"))
        canvas.drawString(2.2 * cm, 1.2 * cm, "Blindspot Lab demonstration unit · synthetic material")
        canvas.drawRightString(A4[0] - 2.2 * cm, 1.2 * cm, f"Page {doc_.page}")
        canvas.restoreState()

    doc.build(s, onFirstPage=footer, onLaterPages=footer)


# ---------------------------------------------------------------------------
# PPTX deck (python-pptx)
# ---------------------------------------------------------------------------


def build_pptx(path: Path, figdir: Path) -> None:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    title_only = prs.slide_layouts[5]
    title_content = prs.slide_layouts[1]

    def notes(slide, text: str) -> None:  # noqa: ANN001
        # The default template is 4:3; stretch the title placeholder to the 16:9 width.
        if slide.shapes.title is not None:
            t = slide.shapes.title
            t.left, t.width = Inches(0.6), Inches(12.1)
            if slide.slide_layout == prs.slide_layouts[0]:
                t.top, t.height = Inches(2.3), Inches(1.5)
            else:
                t.top, t.height = Inches(0.3), Inches(1.1)
        for ph in slide.placeholders:
            if ph.placeholder_format.idx == 1 and slide.slide_layout == prs.slide_layouts[0]:
                ph.left, ph.width, ph.top, ph.height = Inches(0.6), Inches(12.1), Inches(4.0), Inches(1.0)
        slide.notes_slide.notes_text_frame.text = text

    def bullets(title: str, items: list[str], note: str) -> None:
        slide = prs.slides.add_slide(title_content)
        slide.shapes.title.text = title
        body = slide.placeholders[1]
        body.left, body.top, body.width, body.height = Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.3)
        tf = body.text_frame
        tf.text = items[0]
        for it in items[1:]:
            tf.add_paragraph().text = it
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size = Pt(26)
        notes(slide, note)

    # 1 Title
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Unit Economics for Marketers"
    s.placeholders[1].text = "Blindspot Lab demo deck · Lecture 2 · synthetic examples"
    notes(s, "This deck accompanies the primer chapter on contribution margin and ROAS. " + SYNTHETIC_NOTE)

    # 2 Funnel built from native shapes
    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "From Visitors to Customers: The Funnel"
    stages = [("Visitors", "20,000"), ("Product views", "6,000"), ("Add to cart", "1,200"), ("Purchases", "500")]
    fills = ["4F46E5", "4338CA", "0891B2", "059669"]
    for i, (name, n) in enumerate(stages):
        w = Inches(9.0 - i * 1.8)
        left = Emu(int((prs.slide_width - w) / 2))
        shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(1.7 + i * 1.15), w, Inches(0.95))
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(fills[i])
        shp.line.fill.background()
        shp.text_frame.text = f"{name}: {n}"
        for p in shp.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size = Pt(24)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    tb = s.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.5), Inches(0.7))
    tb.text_frame.text = "Overall conversion rate = 500 ÷ 20,000 = 2.5%"
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(22)
    notes(
        s,
        "Each stage of the funnel loses people. The overall conversion rate is purchases divided by visitors, so the "
        "paid search funnel in our example converts at 2.5 percent. The biggest drop is between product views and "
        "add to cart.",
    )

    # 3
    bullets(
        "Revenue Is Not Profit",
        [
            "Revenue = price × units sold",
            "Variable costs rise with every unit sold",
            "Profit appears only after variable and fixed costs are covered",
            "Growing revenue can still shrink profit if each sale earns less",
        ],
        "A common mistake is to celebrate revenue growth that came from discounting. Always ask what each sale "
        "contributes after its variable costs.",
    )
    # 4
    bullets(
        "Contribution Margin",
        [
            "Contribution margin per unit = price − variable cost",
            "Contribution margin ratio = contribution ÷ price",
            "Total contribution = contribution per unit × units sold",
            "Contribution pays for fixed costs first, then becomes profit",
        ],
        "For a ₹1,200 jacket with an ₹800 variable cost, contribution per unit is ₹400 and the contribution margin "
        "ratio is one third.",
    )
    # 5
    bullets(
        "Price Cuts Need More Volume Than You Think",
        [
            "Jacket price: ₹1,200 → ₹1,080 (a 10% cut)",
            "Variable cost stays at ₹800",
            "Contribution per unit: ₹400 → ₹280",
            "Volume needed to keep contribution: 400 ÷ 280 − 1 = +42.9%",
            "Volume needed to keep revenue: only +11.1%",
        ],
        "Students often answer with the revenue figure. The question is what the business keeps after variable "
        "costs, so the contribution figure of 42.9 percent is the one that protects profit.",
    )
    # 6 Picture diagram
    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "Where the Money Goes"
    s.shapes.add_picture(str(figdir / "waterfall.png"), Inches(1.4), Inches(1.5), width=Inches(10.5))
    notes(
        s,
        "The waterfall shows that after the discount most of the price goes to variable costs: only ₹280 of the "
        "₹1,080 discounted price remains as contribution per jacket.",
    )
    # 7 Native chart
    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "ROAS by Channel"
    cd = CategoryChartData()
    cd.categories = ["Search", "Social", "Display", "Email"]
    cd.add_series("ROAS (x)", (4.0, 2.0, 1.5, 9.0))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.8), Inches(1.5), Inches(8.0), Inches(5.3), cd)
    gf.chart.has_legend = False
    gf.chart.plots[0].has_data_labels = True
    tb = s.shapes.add_textbox(Inches(9.1), Inches(2.4), Inches(3.9), Inches(3))
    tb.text_frame.word_wrap = True
    tb.text_frame.text = "Break-even ROAS at a 40% contribution margin ratio = 2.5"
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(22)
    notes(
        s,
        "Display and Social are above 1 but below the break-even ROAS of 2.5, so they lose money on a contribution "
        "basis. Search and Email are above break-even.",
    )
    # 8 Native table
    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "Break-even ROAS"
    rows = [("Contribution margin ratio", "Break-even ROAS"), ("20%", "5.0x"), ("25%", "4.0x"), ("40%", "2.5x"),
            ("50%", "2.0x"), ("60%", "1.67x")]
    tbl = s.shapes.add_table(len(rows), 2, Inches(2.5), Inches(1.6), Inches(8.3), Inches(4.6)).table
    for r, (a, b) in enumerate(rows):
        tbl.cell(r, 0).text = a
        tbl.cell(r, 1).text = b
        for c in (0, 1):
            for p in tbl.cell(r, c).text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(22)
    notes(s, "Break-even ROAS is one divided by the contribution margin ratio. Lower margins need much higher ROAS.")
    # 9
    bullets(
        "Acquisition Economics",
        [
            "CAC = sales and marketing spend ÷ new customers",
            "CLV should be based on contribution, not revenue",
            "Payback period = CAC ÷ monthly contribution per customer",
            "Example: ₹600 monthly revenue at a 30% margin = ₹180 contribution per month",
        ],
        "A subscription box customer paying ₹600 a month at a 30 percent contribution margin ratio generates ₹180 of "
        "contribution per month. With a CAC of ₹720 the payback period is four months.",
    )
    # 10
    bullets(
        "Correlation Is Not Causation",
        [
            "Loyalty-app users spend more, but who chooses to install the app?",
            "Confounders can create correlations without causation",
            "Randomised A/B tests balance confounders on average",
            "Ask whether an effect is large enough to matter, not only whether it is significant",
        ],
        "In our synthetic test, the holdout group that never saw the loyalty-app promotion spent ₹1,950 per month on "
        "average, compared with ₹2,050 for the promoted group, a lift of about 5 percent.",
    )
    # 11
    bullets(
        "Key Takeaways",
        [
            "Judge pricing and campaigns on contribution, not revenue",
            "Break-even ROAS = 1 ÷ contribution margin ratio",
            "Compare CAC with contribution-based CLV",
            "Use experiments before claiming causal effects",
        ],
        SYNTHETIC_NOTE,
    )
    prs.core_properties.title = "Unit Economics for Marketers"
    prs.core_properties.author = "Blindspot Lab team (synthetic demo material)"
    prs.save(str(path))


# ---------------------------------------------------------------------------
# Narrated lecture video (Piper TTS + ffmpeg)
# ---------------------------------------------------------------------------

LECTURE = [
    ("Lecture 3\nWhy Price Cuts Need More Volume\nThan You Think",
     "Welcome to lecture three of the Blindspot Lab demo unit. This short lecture uses synthetic examples to explain "
     "why a price cut usually needs much more extra volume than people expect. The narration is computer generated."),
    ("Contribution margin\n= price - variable cost",
     "Start with contribution margin. Contribution margin per unit is the selling price minus the variable cost of "
     "producing and delivering one unit. It is the money each sale contributes towards fixed costs and profit."),
    ("Coffee pods\nPrice ₹40 · Variable cost ₹25\nContribution ₹15 per pod",
     "Imagine a coffee pod sold for forty rupees, with a variable cost of twenty five rupees. Each pod contributes "
     "fifteen rupees."),
    ("Cut price to ₹35\nContribution falls to ₹10 per pod",
     "Now the marketing team proposes cutting the price to thirty five rupees. The variable cost does not change, so "
     "the contribution per pod falls from fifteen rupees to ten rupees. A twelve and a half percent price cut has "
     "removed one third of the contribution per unit."),
    ("Preserve contribution\n15 ÷ 10 - 1 = +50% volume",
     "To earn the same total contribution, the business must sell fifteen divided by ten, which is one and a half "
     "times as many pods. That is a fifty percent increase in volume."),
    ("Preserve revenue\n40 ÷ 35 - 1 = +14.3% volume",
     "If you only try to keep revenue constant, you would aim for about fourteen percent more volume. That target "
     "ignores variable costs, so hitting it would still leave the business with less contribution than before."),
    ("Revenue preservation\nis not\ncontribution preservation",
     "So always ask which quantity you are trying to preserve. Revenue preservation and contribution preservation are "
     "different goals, and for pricing decisions, contribution is usually the one that matters."),
    ("Break-even ROAS\n= 1 ÷ contribution margin ratio",
     "The same logic applies to advertising. A return on advertising spend above one does not mean a campaign is "
     "profitable. With a contribution margin ratio of thirty seven point five percent, as in our coffee example, the "
     "break-even ROAS is about two point seven."),
    ("Correlation is not causation\nUse a randomised control group",
     "Finally, be careful with evidence. If customers who received a discount bought more, that correlation does not "
     "prove the discount caused the extra purchases. A randomised experiment with a control group is the cleanest way "
     "to measure the causal effect."),
    ("Summary\nCompare contribution, not revenue\nCheck break-even ROAS\nLook for experimental evidence",
     "To summarise: compare contribution, not revenue. Compute the break-even ROAS before celebrating a campaign. "
     "And look for experimental evidence before claiming that a marketing action caused a result. Thanks for "
     "watching."),
]


def render_frame(text: str, path: Path, index: int) -> None:
    w, h = 1280, 720
    img = Image.new("RGB", (w, h), (15, 23, 42))
    d = ImageDraw.Draw(img)
    for y in range(h):  # subtle vertical gradient
        c = int(15 + 18 * y / h)
        d.line([(0, y), (w, y)], fill=(c, c + 8, 42 + int(20 * y / h)))
    big = ImageFont.truetype(str(FONT_DIR / "DejaVuSans-Bold.ttf"), 54)
    small = ImageFont.truetype(str(FONT_DIR / "DejaVuSans.ttf"), 22)
    lines = text.split("\n")
    total = len(lines) * 74
    y = (h - total) // 2
    for i, line in enumerate(lines):
        font = big
        tw = d.textlength(line, font=font)
        d.text(((w - tw) / 2, y), line, font=font, fill=(224, 231, 255) if i == 0 else (165, 243, 252))
        y += 74
    d.text((40, h - 50), f"Blindspot Lab · Lecture 3 · synthetic narration · {index + 1}/{len(LECTURE)}",
           font=small, fill=(148, 163, 184))
    img.save(path)


def wav_duration(path: Path) -> float:
    with wave.open(str(path)) as wf:
        return wf.getnframes() / wf.getframerate()


def build_video(path: Path, voice: Path, work: Path) -> list[dict]:
    segments = []
    concat_lines = []
    t = 0.0
    for i, (frame_text, narration) in enumerate(LECTURE):
        frame = work / f"frame_{i:02d}.png"
        render_frame(frame_text, frame, i)
        wav = work / f"seg_{i:02d}.wav"
        subprocess.run(
            ["python", "-m", "piper", "-m", str(voice), "-f", str(wav), "--sentence-silence", "0.35"],
            input=narration.encode(), check=True, capture_output=True,
        )
        dur = wav_duration(wav) + 0.6  # small pause between segments
        clip = work / f"clip_{i:02d}.mp4"
        subprocess.run(
            [
                "ffmpeg", "-y", "-loop", "1", "-i", str(frame), "-i", str(wav), "-af", "apad=pad_dur=0.6",
                "-c:v", "libx264", "-tune", "stillimage", "-pix_fmt", "yuv420p", "-r", "25",
                "-c:a", "aac", "-b:a", "128k", "-ar", "44100", "-t", f"{dur:.3f}", str(clip),
            ],
            check=True, capture_output=True,
        )
        concat_lines.append(f"file '{clip}'")
        segments.append({"index": i, "start": round(t, 2), "end": round(t + dur, 2),
                         "frame_text": frame_text, "narration": narration})
        t += dur
    lst = work / "concat.txt"
    lst.write_text("\n".join(concat_lines))
    subprocess.run(
        ["ffmpeg", "-y", "-f", "concat", "-safe", "0", "-i", str(lst), "-c", "copy", "-movflags", "+faststart",
         "-metadata", "title=Lecture 3: Why Price Cuts Need More Volume Than You Think (synthetic narration)",
         str(path)],
        check=True, capture_output=True,
    )
    return segments


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--voice", type=Path, default=Path.home() / ".cache/piper/en_US-lessac-medium.onnx")
    ap.add_argument("--skip-video", action="store_true")
    args = ap.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        fig_price_cut(work / "fig2_1.png")
        fig_roas_channels(work / "fig3_1.png")
        fig_order_values(work / "fig6_1.png")
        fig_waterfall(work / "waterfall.png")
        build_pdf(OUT / "Marketing_Analytics_Primer.pdf", work)
        build_pptx(OUT / "Unit_Economics_for_Marketers.pptx", work)
        print("PDF and PPTX written")
        if not args.skip_video:
            segs = build_video(OUT / "Lecture_03_Price_Cuts.mp4", args.voice, work)
            (OUT / "lecture_03_narration.json").write_text(
                json.dumps(
                    {
                        "title": "Lecture 3: Why Price Cuts Need More Volume Than You Think",
                        "voice": "Piper TTS en_US-lessac-medium (synthesized speech)",
                        "note": "Planned segment timings; ground truth for evaluation only. The application "
                        "derives its own timestamps by transcribing the MP4.",
                        "segments": segs,
                    },
                    indent=2,
                    ensure_ascii=False,
                )
            )
            print(f"Video written ({segs[-1]['end']:.1f}s)")


if __name__ == "__main__":
    main()
